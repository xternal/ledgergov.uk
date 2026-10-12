"""
OBR's assessment of the government's fiscal targets (the "headroom"), at every EFO that made one.

The OBR judges the fiscal mandate and supplementary targets in the Charter for Budget
Responsibility in one table of each assessing Economic and fiscal outlook (EFO): "Table 7.1:
Performance against the Government's fiscal targets" (October 2024: "... proposed fiscal
targets", the rules as proposed then and approved by Parliament in January 2025). Since the
Charter for Budget Responsibility: Autumn 2025 (approved 24 February 2026) the OBR assesses
the rules once a year, at the Budget, so a spring EFO has no such table (March 2026 says so in
its foreword). This module records each assessing edition's own figures, one file per edition.

Editions (fetch):
  * every edition in ASSESSED (the assessments under the current rules) that has no committed
    file yet, and the newest of them every time, so SOURCE describes it;
  * the newest EFO named on https://obr.uk/efo/ if it is newer than all of those and not in
    NOT_ASSESSED: its last "charts and tables: Chapter N" workbook is opened and used if it holds
    the table (on Budget day the nightly picks the new assessment up this way). If it does not,
    that is logged and nothing is recorded for it; add it to NOT_ASSESSED to stop asking.
  Each edition's page gives its title and publication date and links the chapter workbooks; the
  fiscal targets chapter has been the last one in every assessing edition so far.

OBR refuses GitHub's servers (403); every request may go through the Internet Archive
(core.download(archive=True)). A failure is not caught here, so the build keeps the committed files.

Table read (found by its title, read by labels, never by cell position):

  Sections, each naming the target year:
    "Current budget|balance to be in surplus by|in <year>"            -> rule current_budget (the stability rule)
    "Change in public sector net financial liabilities in <year>"     -> rule psnfl (the investment rule)
    "Welfare cap: specified welfare spending in <year>"               -> rule welfare_cap
  Rows: "<Month Year> forecast" for the edition itself, and "<Month Year> pre-measures forecast"
  (the forecast before that fiscal event's policy decisions). Rows restating an earlier edition
  and "Memo:" rows are not recorded: each edition is read from its own table.
  Columns, by their two header rows: "Per cent of GDP"/"£ billion" x "Forecast"/"Margin", "Per cent"/"Probability".

Series (period = the target year, kind "forecast", unit as named):
  fiscal_targets.<rule>.margin                  gbp_bn   margin by which the target is met (negative: missed)
  fiscal_targets.<rule>.margin_pct_gdp          pct_gdp  the same in per cent of GDP (not for the welfare cap)
  fiscal_targets.<rule>.probability             pct      OBR's probability of meeting the target (when given)
  fiscal_targets.<rule>.margin_pre_measures     gbp_bn   the margin before that event's policy decisions (when given)
  fiscal_targets.welfare_cap.spending           gbp_bn   forecast spending within the cap
Each observation's method_note names the table and row it was read from.
"""

from __future__ import annotations

import calendar
import logging
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path

import httpx
import openpyxl

from etl.core import BUILD_DIR, Observation, RawArtifact, Source, download, links
from etl.sources.obr_efo import _canonical, _strip_query, edition_candidates, edition_of, page_published_on, page_title

log = logging.getLogger(__name__)

INDEX_URL = "https://obr.uk/efo/"  # redirects to the newest EFO

# Every EFO that assessed the current fiscal rules, oldest first: vintage -> edition page.
ASSESSED: dict[str, str] = {
    "EFO-2024-10": "https://obr.uk/efo/economic-and-fiscal-outlook-october-2024/",
    "EFO-2025-03": "https://obr.uk/efo/economic-and-fiscal-outlook-march-2025/",
    "EFO-2025-11": "https://obr.uk/efo/economic-and-fiscal-outlook-november-2025/",
}
# EFOs known to make no assessment (a spring forecast since the Charter for Budget Responsibility:
# Autumn 2025): discovery does not open their workbooks.
NOT_ASSESSED: set[str] = {"EFO-2026-03"}

SOURCE = Source(
    id="obr_fiscal_targets",
    title="OBR Economic and fiscal outlook – November 2025: performance against the fiscal targets",
    publisher="Office for Budget Responsibility",
    url=ASSESSED["EFO-2025-11"],
    licence="OGL v3",
    published_on=date(2025, 11, 26),
    # One assessment a year, at the Budget (late October or November). fetch() sets published_on from
    # the newest assessing edition, so staleness counts from it.
    cadence_days=365,
    grace_days=45,
)

CHAPTER_WORKBOOK = r"/download/[^/?#]*charts-and-tables-chapter-(\d+)/"
TABLE_TITLE = re.compile(r"^Table\s+(\d+\.\d+):\s*(Performance against the Government['’]s (?:proposed )?fiscal targets)\s*$", re.I)

SECTIONS = {
    "current_budget": re.compile(r"^current (?:budget|balance) to be in surplus (?:by|in) (\d{4}-\d{2})$", re.I),
    "psnfl": re.compile(r"^change in public sector net financial liabilities in (\d{4}-\d{2})$", re.I),
    "welfare_cap": re.compile(r"^welfare cap: specified welfare spending in (\d{4}-\d{2})$", re.I),
}
ROW = re.compile(r"^([A-Z][a-z]+ \d{4}) (pre-measures )?forecast$")

_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}
_FILE = re.compile(r"^efo_(\d{4})_(\d{2})_fiscal_targets\.xlsx$")


@dataclass
class Edition:
    url: str
    year: int
    month: int

    @property
    def vintage(self) -> str:
        return f"EFO-{self.year}-{self.month:02d}"

    @property
    def label(self) -> str:
        return f"{calendar.month_name[self.month]} {self.year}"

    @property
    def prefix(self) -> str:
        return f"efo_{self.year}_{self.month:02d}"


def edition(url: str) -> Edition:
    ym = edition_of(url)
    if ym is None:
        raise LookupError(f"{url} is not an EFO edition page")
    return Edition(url, *ym)


def committed_vintages() -> set[str]:
    return {p.stem for p in (BUILD_DIR / "observations" / SOURCE.id).glob("EFO-*.csv")}


def _page(ed: Edition) -> str:
    return download(SOURCE.id, ed.url, f"{ed.prefix}_landing.html", archive=True).path.read_text(errors="replace")


def last_chapter_workbook(url: str, html: str) -> str:
    """The highest-numbered 'charts and tables: Chapter N' workbook linked on an edition page."""
    found: dict[int, str] = {}
    for u in links(html, url, CHAPTER_WORKBOOK):
        m = re.search(CHAPTER_WORKBOOK, u)
        found.setdefault(int(m.group(1)), _strip_query(u))
    if not found:
        raise LookupError(f"{url}: no charts-and-tables chapter workbooks")
    return found[max(found)]


def _workbook(ed: Edition, html: str) -> RawArtifact:
    return download(SOURCE.id, last_chapter_workbook(ed.url, html), f"{ed.prefix}_fiscal_targets.xlsx", vintage=ed.vintage, archive=True)


def describe(ed: Edition, html: str) -> None:
    title = page_title(html) or f"OBR Economic and fiscal outlook – {ed.label}"
    SOURCE.url, SOURCE.title, SOURCE.published_on = ed.url, f"{title}: performance against the fiscal targets", page_published_on(html)


def newest_listed() -> Edition | None:
    """The newest EFO named on obr.uk/efo/, or None if the page cannot be read."""
    try:
        html = download(SOURCE.id, INDEX_URL, "efo_index.html", archive=True).path.read_text(errors="replace")
    except (httpx.HTTPError, OSError) as e:
        log.warning("obr_fiscal_targets: cannot read %s (%s); checking the known assessments only", INDEX_URL, e)
        return None
    candidates = edition_candidates(html, _canonical(html, INDEX_URL) or INDEX_URL)
    return edition(candidates[0]) if candidates else None


def fetch(since: date | None = None) -> list[RawArtifact]:
    """
    The workbooks of every assessing edition not yet committed, and of the newest one; SOURCE then
    describes the newest. `since` is accepted for the common contract and not used.
    """
    known = [edition(u) for u in ASSESSED.values()]
    have = committed_vintages()
    newest = max(known, key=lambda e: (e.year, e.month))
    listed = newest_listed()
    if listed and (listed.year, listed.month) > (newest.year, newest.month) and listed.vintage not in NOT_ASSESSED and listed.vintage not in ASSESSED:
        html = _page(listed)
        raw = _workbook(listed, html)
        if has_table(raw.path):
            log.info("obr_fiscal_targets: %s assesses the fiscal rules; recording it", listed.vintage)
            known.append(listed)
            newest = listed
        else:
            log.info("obr_fiscal_targets: %s has no fiscal targets table (a forecast without an assessment); "
                     "the newest assessment stays %s. Add it to NOT_ASSESSED to stop checking.", listed.vintage, newest.vintage)
    out = []
    for ed in known:
        if ed.vintage in have and ed is not newest:
            continue
        html = _page(ed)
        out.append(_workbook(ed, html))
        if ed is newest:
            describe(ed, html)
    log.info("obr_fiscal_targets: newest assessment %s (%s, published %s)", newest.vintage, SOURCE.title, SOURCE.published_on)
    return out


# --------------------------------------------------------------------------- parsing


@dataclass
class TableRow:
    rule: str
    target_year: str
    label: str  # e.g. "November 2025 forecast"
    edition: str  # "November 2025"
    pre_measures: bool
    values: dict[tuple[str, str], float]  # (unit group, column) -> value, e.g. ("gbp_bn", "margin")


@dataclass
class FiscalTable:
    number: str  # "7.1"
    title: str
    rows: list[TableRow]

    @property
    def name(self) -> str:
        return f"Table {self.number}"


def _text(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if isinstance(v, str) else ""


def _clean(label: str) -> str:
    """'Probability1' -> 'Probability'; footnote digits after a word dropped."""
    return re.sub(r"(?<=[A-Za-z])\d+$", "", label).strip()


def find_sheet(wb):
    """(sheet, table number, title) of the fiscal targets table, or None. The Contents sheet lists the title too."""
    for ws in wb.worksheets:
        if ws.title.strip().lower() == "contents":
            continue
        for row in ws.iter_rows(max_row=4, values_only=True):
            for v in row:
                if m := TABLE_TITLE.match(_text(v)):
                    return ws, m.group(1), m.group(2).replace("’", "'")  # one apostrophe in every edition's title
    return None


def has_table(path: Path) -> bool:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return find_sheet(wb) is not None
    finally:
        wb.close()


GROUPS = {"per cent of gdp": "pct_gdp", "£ billion": "gbp_bn", "per cent": "pct"}
COLUMNS = {"forecast": "forecast", "margin": "margin", "probability": "probability"}


def read_table(path: Path) -> FiscalTable:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    found = find_sheet(wb)
    if found is None:
        raise ValueError(f"{path.name}: no 'Performance against the Government's fiscal targets' table")
    ws, number, title = found
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    # Two header rows: unit groups ("Per cent of GDP", "£ billion", "Per cent"), then columns under them.
    group_row = next((i for i, r in enumerate(grid) if sum(_text(v).lower() in GROUPS for v in r) >= 2), None)
    if group_row is None or group_row + 1 >= len(grid):
        raise ValueError(f"{path.name} {ws.title}: no unit header row")
    groups = sorted((j, GROUPS[_text(v).lower()]) for j, v in enumerate(grid[group_row]) if _text(v).lower() in GROUPS)
    columns: dict[int, tuple[str, str]] = {}
    for j, v in enumerate(grid[group_row + 1]):
        col = COLUMNS.get(_clean(_text(v)).lower())
        if col is None:
            continue
        group = [g for k, g in groups if k <= j]
        if not group:
            raise ValueError(f"{path.name} {ws.title}: column {col!r} has no unit above it")
        columns[j] = (group[-1], col)
    if ("gbp_bn", "margin") not in columns.values():
        raise ValueError(f"{path.name} {ws.title}: no '£ billion' 'Margin' column")
    first_value_col = min(columns)

    rows: list[TableRow] = []
    section: tuple[str, str] | None = None
    for r in grid[group_row + 2 :]:
        label = next((_text(v) for j, v in enumerate(r) if j < first_value_col and _text(v)), "")
        if not label:
            continue
        hit = next(((rule, m.group(1)) for rule, rx in SECTIONS.items() if (m := rx.match(label))), None)
        if hit:
            section = hit
            continue
        m = ROW.match(label)
        if not m or section is None:
            continue  # memo rows, footnotes, source lines
        values = {columns[j]: float(v) for j, v in enumerate(r) if j in columns and isinstance(v, (int, float)) and not isinstance(v, bool)}
        rows.append(TableRow(section[0], section[1], label, m.group(1), bool(m.group(2)), values))
    seen_rules = {row.rule for row in rows}
    missing = [rule for rule in ("current_budget", "psnfl") if rule not in seen_rules]
    if missing:
        raise ValueError(f"{path.name} {ws.title}: no rows for {missing}")
    return FiscalTable(number, title, rows)


def edition_from_file(path: Path) -> Edition:
    m = _FILE.match(path.name)
    if not m:
        raise ValueError(f"unexpected file name {path.name}; expected efo_<year>_<month>_fiscal_targets.xlsx")
    year, month = int(m.group(1)), int(m.group(2))
    return Edition(f"https://obr.uk/efo/economic-and-fiscal-outlook-{calendar.month_name[month].lower()}-{year}/", year, month)


def _obs(rule: str, what: str, row: TableRow, value: float, unit: str, vintage: str, table: FiscalTable) -> Observation:
    return Observation(
        series_id=f"fiscal_targets.{rule}.{what}",
        period=row.target_year,
        value=round(value, 6),
        unit=unit,
        kind="forecast",
        source_id=SOURCE.id,
        vintage=vintage,
        quality="sourced",
        method_note=f"{table.name} ({table.title}), row '{row.label}'.",
    )


def parse(raws: list[RawArtifact]) -> list[Observation]:
    out: list[Observation] = []
    for raw in raws if isinstance(raws, list) else [raws]:
        ed = edition_from_file(raw.path)
        vintage = raw.vintage or ed.vintage
        if vintage != ed.vintage:
            raise ValueError(f"{raw.path.name}: file is {ed.vintage} but the artifact says {vintage}")
        table = read_table(raw.path)
        own = [r for r in table.rows if r.edition == ed.label]
        if not any(r.rule == "current_budget" and not r.pre_measures for r in own):
            raise ValueError(f"{raw.path.name}: {table.name} has no '{ed.label} forecast' row for the current budget")
        for row in own:
            v = row.values
            margin = v.get(("gbp_bn", "margin"))
            if margin is None:
                raise ValueError(f"{raw.path.name}: row '{row.label}' ({row.rule}) has no £ billion margin")
            if row.pre_measures:
                out.append(_obs(row.rule, "margin_pre_measures", row, margin, "gbp_bn", vintage, table))
                continue
            out.append(_obs(row.rule, "margin", row, margin, "gbp_bn", vintage, table))
            if ("pct_gdp", "margin") in v:
                out.append(_obs(row.rule, "margin_pct_gdp", row, v[("pct_gdp", "margin")], "pct_gdp", vintage, table))
            if ("pct", "probability") in v:
                out.append(_obs(row.rule, "probability", row, v[("pct", "probability")], "pct", vintage, table))
            if row.rule == "welfare_cap" and ("gbp_bn", "forecast") in v:
                out.append(_obs(row.rule, "spending", row, v[("gbp_bn", "forecast")], "gbp_bn", vintage, table))
    seen = set()
    for o in out:
        k = (o.series_id, o.period, o.vintage)
        if k in seen:
            raise ValueError(f"obr_fiscal_targets: duplicate observation {k}")
        seen.add(k)
    return out
